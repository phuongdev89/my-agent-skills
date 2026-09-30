# -*- coding: utf-8 -*-
"""
verify_docx.py
Kịch bản tự động kiểm tra chất lượng file DOCX sau khi biên dịch từ PDF/ảnh:
1. Đếm số đoạn (Paragraphs), số bảng (Tables).
2. Đếm số lần ngắt trang tường minh (<w:br w:type="page"/>) và so sánh với số trang kỳ vọng.
3. Kiểm tra danh sách Heading 1 & Heading 2 (đảm bảo hiển thị chuẩn trên Left Navigation Pane).
4. Tầm soát rò rỉ header/footer (dò các đoạn văn chứa tiêu đề lặp, số trang đơn lẻ).
5. Tổng hợp cấu trúc kích thước các bảng biểu.

Cách dùng:
    python verify_docx.py <path_to_docx> [expected_pages]
Ví dụ:
    python verify_docx.py "output.docx" 36
"""

import os
import sys
import re
import docx

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

def verify_document(doc_path, expected_pages=None):
    if not os.path.isfile(doc_path):
        print(f"Lỗi: Không tìm thấy tệp DOCX tại: {doc_path}")
        sys.exit(1)

    print("=" * 75)
    print(f"BÁO CÁO NGHIỆM THU TỆP DOCX: {os.path.basename(doc_path)}")
    print("=" * 75)

    doc = docx.Document(doc_path)
    total_paras = len(doc.paragraphs)
    total_tables = len(doc.tables)

    # 1. Đếm số ngắt trang tường minh
    page_breaks = 0
    for p in doc.paragraphs:
        for r in p.runs:
            if '<w:br' in r._r.xml and 'type="page"' in r._r.xml:
                page_breaks += 1

    actual_pages = page_breaks + 1

    print(f"• Tổng số đoạn văn (Paragraphs): {total_paras}")
    print(f"• Tổng số bảng biểu (Tables): {total_tables}")
    print(f"• Số ngắt trang (Page breaks): {page_breaks}")
    print(f"• Tổng số trang ước tính: {actual_pages}")

    if expected_pages is not None:
        if actual_pages == expected_pages:
            print(f"  ==> [ĐẠT CHUẨN 1:1] Số trang khớp hoàn toàn với bản gốc: {actual_pages}/{expected_pages}")
        else:
            print(f"  ==> [CẢNH BÁO LỆCH TRANG] Kỳ vọng {expected_pages} trang, tài liệu hiện có {actual_pages} trang (Lệch {actual_pages - expected_pages:+d})")

    # 2. Kiểm tra Navigation Tree (Headings)
    h1s = [p.text.strip() for p in doc.paragraphs if p.style.name == 'Heading 1' and p.text.strip()]
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if p.style.name == 'Heading 1' and p.text.strip():
                        if p.text.strip() not in h1s:
                            h1s.append(p.text.strip())
    h2s = []
    # Quét Heading 2 trong cả paragraph ngoài và trong bảng banner
    for p in doc.paragraphs:
        if p.style.name == 'Heading 2' and p.text.strip():
            h2s.append(p.text.strip())
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if p.style.name == 'Heading 2' and p.text.strip():
                        if p.text.strip() not in h2s:
                            h2s.append(p.text.strip())

    print(f"\n• Cây Điều Hướng (Navigation Pane):")
    print(f"  - Heading 1 (Chương / Phần lớn): {len(h1s)}")
    for i, h in enumerate(h1s, 1):
        print(f"    {i}. {h}")
    print(f"  - Heading 2 (Đề mục banner): {len(h2s)}")
    for i, h in enumerate(h2s[:10], 1): # In 10 cái đầu
        print(f"    {i}. {h}")
    if len(h2s) > 10:
        print(f"    ... và {len(h2s) - 10} mục khác.")

    # 3. Tầm soát rò rỉ Header / Footer
    suspicious_leaks = []
    for p in doc.paragraphs:
        t = p.text.strip()
        # Chữ số đứng 1 mình (số trang bị trôi vào body)
        if t.isdigit() and len(t) <= 4:
            suspicious_leaks.append(f"Số trang trôi nổi: '{t}'")
        # Chuỗi running header thường gặp bị copy vào body
        if re.search(r'^(CHƯƠNG \d+|PHẦN \d+|GIÁO TRÌNH|BÀI \d+|MỤC \d+)\b', t, re.IGNORECASE) and len(t) < 40:
            if p.style.name not in ['Heading 1', 'Heading 2']:
                suspicious_leaks.append(f"Dòng nghi vấn header: '{t}'")

    print(f"\n• Kiểm tra rò rỉ Header/Footer chạy chân trang:")
    if not suspicious_leaks:
        print("  ==> [AN TOÀN] Không phát hiện văn bản running header hoặc số trang rò rỉ vào thân bài.")
    else:
        print(f"  ==> [CẢNH BÁO] Phát hiện {len(suspicious_leaks)} dòng nghi vấn:")
        for leak in suspicious_leaks[:5]:
            print(f"      ! {leak}")

    # 4. Thống kê bảng
    print(f"\n• Danh sách bảng biểu đã vẽ lại:")
    for idx, tbl in enumerate(doc.tables[:8], 1):
        rows = len(tbl.rows)
        cols = len(tbl.columns)
        first_txt = tbl.rows[0].cells[0].text.strip()[:35].replace('\n', ' ')
        print(f"  Bảng {idx:2d}: {rows} hàng x {cols} cột | Khởi đầu: '{first_txt}'")
    if len(doc.tables) > 8:
        print(f"  ... và {len(doc.tables) - 8} bảng biểu khác.")

    print("\n" + "=" * 75)
    print("HOÀN TẤT KIỂM ĐỊNH!")
    print("=" * 75)

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Cách dùng: python verify_docx.py <path_to_docx> [expected_pages]")
        sys.exit(1)

    d_path = sys.argv[1]
    exp_pages = int(sys.argv[2]) if len(sys.argv) > 2 else None
    verify_document(d_path, exp_pages)
