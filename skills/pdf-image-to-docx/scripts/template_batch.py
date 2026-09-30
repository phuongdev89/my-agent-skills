# -*- coding: utf-8 -*-
"""
template_batch.py
Mẫu khung chuẩn để viết một file batch (10-12 trang).
Quy tắc vàng:
1. Mỗi trang trong PDF tương ứng 1 khối code rõ ràng, có comment chỉ số trang (# --- TRANG X ---).
2. Cuối mỗi trang (trừ trang cuối cùng của tài liệu), gọi add_explicit_page_break(doc).
3. Đề mục lớn dùng add_chapter_title() hoặc add_heading_2_banner().
4. Bảng biểu, sơ đồ vẽ bằng docx_helpers (add_comparison_cards, add_process_steps, add_stat_boxes, v.v.).
"""

import os
import sys
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from pathlib import Path

# Import helpers từ cùng thư mục scripts, hoặc từ thư mục gốc của skill khi đặt trong session
try:
    import docx_helpers as h
except ImportError:
    _current_dir = Path(__file__).resolve().parent
    if (_current_dir / "docx_helpers.py").exists():
        sys.path.insert(0, str(_current_dir))
    for _parent in _current_dir.parents:
        _skill_scripts = _parent / "skills" / "pdf-image-to-docx" / "scripts"
        if _skill_scripts.exists():
            sys.path.insert(0, str(_skill_scripts))
            break
        _agent_skill_scripts = _parent / ".agents" / "skills" / "pdf-image-to-docx" / "scripts"
        if _agent_skill_scripts.exists():
            sys.path.insert(0, str(_agent_skill_scripts))
            break
    import docx_helpers as h

def add_batch_1(doc):
    print("  -> Đang nạp Batch 1 (Trang 1 đến 12)...")

    # =========================================================================
    # TRANG 1 (Tương ứng file page_1.png hoặc trang 1 của PDF)
    # =========================================================================
    h.add_chapter_title(doc, "CHƯƠNG 1", "TÊN CHƯƠNG MẪU Ở ĐÂY")

    h.add_heading_2_banner(doc, "1. TIÊU ĐỀ ĐỀ MỤC MẪU BANNER")

    h.add_body_p(doc, [
        ("Đây là đoạn văn bản mẫu giữ nguyên 100% từng câu chữ từ bản gốc. ", False, False, None),
        ("Từ khóa quan trọng được in đậm", True, False, RGBColor(15, 23, 42)),
        (", tiếp tục các câu từ tiếp theo theo đúng bản scan.", False, False, None)
    ])

    h.add_bullet_p(doc, [
        ("Ý chính thứ nhất: ", True, False, None),
        ("Nội dung diễn giải chi tiết...", False, False, None)
    ])

    # Kết thúc trang 1 -> ngắt trang sang trang 2
    h.add_explicit_page_break(doc)

    # =========================================================================
    # TRANG 2 (Tương ứng file page_2.png)
    # =========================================================================
    h.add_heading_2_banner(doc, "2. SƠ ĐỒ HOẶC BẢNG SO SÁNH TRỰC QUAN")

    # Ví dụ vẽ lại sơ đồ so sánh 2 cột thay vì chụp ảnh
    h.add_comparison_cards(
        doc,
        left_card={
            'badge': '❌ CÁCH CŨ / SAI',
            'badge_color': '991B1B',
            'bg_hex': 'FEF2F2',
            'border_color': 'FECACA',
            'title': 'Làm thủ công tốn thời gian',
            'points': ['Mất 4-5 tiếng mỗi video', 'Chất lượng không đồng đều', 'Dễ nản lòng']
        },
        right_card={
            'badge': '✔ CÁCH MỚI / ĐÚNG',
            'badge_color': '166534',
            'bg_hex': 'F0FDF4',
            'border_color': 'BBF7D0',
            'title': 'Ứng dụng AI & Batching theo lô',
            'points': ['Chỉ 15-20 phút mỗi video', 'Chất lượng chuẩn hóa cao', 'Dễ dàng scale 5-10 kênh']
        }
    )

    # Kết thúc trang 2
    h.add_explicit_page_break(doc)

    # Tiếp tục cho đến hết các trang trong Batch...
    print("  ✓ Hoàn thành Batch 1!")
